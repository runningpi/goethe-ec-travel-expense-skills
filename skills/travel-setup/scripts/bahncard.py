"""BahnCard reimbursement: savings table (xlsx + pdf) and request letter (pdf).

  python bahncard.py --data bahncard.json --outdir <folder>

bahncard.json:
{
  "type": 50,                          # 25 or 50 (default: profile.bahncard.type)
  "number": "1234 5678 9012 3456",  # default: profile
  "price": 79.90, "purchase_date": "10.08.2026",
  "cost_center": "3502030005",
  "date": "06.10.2026",                # letter date, default today
  "trips": [
    {"route": "Frankfurt - Bonn - Frankfurt (Hin- und Rückfahrt)", "date": "11.08.2026",
     "paid": 46.45, "fare": "Super Sparpreis Young"}
  ]
}
Discount rule (same as the chair's Excel template): BahnCard 50 gives 50 % on Flexpreis /
Normalpreis / Flexpreis Europa and 25 % on all Sparpreis fares; BahnCard 25 gives 25 %.
Price without BahnCard = paid / (1 - discount); saving = difference.
Writes 02_Uebersicht_Fahrten_BahnCard.xlsx/.pdf and 01_Erklaerung_Erstattung_BahnCard.pdf.
The letter is NOT signed - the user signs it.
"""
import argparse
import datetime
import json
import os

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import ensure_utf8_stdout, load_profile, money, parse_money

FULL_FARES = ("flexpreis", "normalpreis", "flexpreis europa")


def discount(bc_type, fare):
    if int(bc_type) == 50 and fare.strip().lower() in FULL_FARES:
        return 0.5
    return 0.25


def main():
    ensure_utf8_stdout()
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--outdir", required=True)
    a = ap.parse_args()
    d = json.load(open(a.data, encoding="utf-8"))
    prof = load_profile()
    bc = prof.get("bahncard", {})
    bc_type = int(str(d.get("type") or bc.get("type") or "50").split()[-1])
    number = d.get("number") or bc.get("number", "")
    price = parse_money(d.get("price") or bc.get("price"))
    date = d.get("date") or datetime.date.today().strftime("%d.%m.%Y")
    os.makedirs(a.outdir, exist_ok=True)

    rows = []
    for t in d["trips"]:
        paid = parse_money(t["paid"])
        disc = discount(bc_type, t["fare"])
        full = paid / (1 - disc)
        rows.append([t["route"], t["date"], paid, t["fare"], disc, full, full - paid])
    saving = round(sum(r[6] for r in rows), 2)  # round once at the end, like Excel

    # --- xlsx (values + formulas, same layout as the chair template)
    wb = Workbook()
    ws = wb.active
    ws.title = "BahnCard"
    ws.append(["BahnCard Nummer", number])
    ws.append(["BahnCard Typ (25/50)", bc_type])
    ws.append(["Kosten BahnCard", price])
    ws.append(["Ersparnis BahnCard", f"=SUM(G8:G{7 + len(rows)})"])
    ws.append([])
    ws.append([])
    ws.append(["Fahrt", "Datum", "Kosten (inkl. BahnCard Rabatt)", "Tarif", "BahnCard Rabatt",
               "Preis ohne BahnCard", "Ersparnis"])
    for i, r in enumerate(rows, start=8):
        ws.append([r[0], r[1], r[2], r[3], r[4], f"=C{i}/(1-E{i})", f"=F{i}-C{i}"])
        ws[f"E{i}"].number_format = "0%"
        for c in "CFG":
            ws[f"{c}{i}"].number_format = '#,##0.00 "€"'
    for c in ("B3", "B4"):
        ws[c].number_format = '#,##0.00 "€"'
    for c in ("A1", "A2", "A3", "A4"):
        ws[c].font = Font(bold=True)
    tab = Table(displayName="Fahrten", ref=f"A7:G{7 + len(rows)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    for col, w in zip("ABCDEFG", (55, 12, 28, 24, 16, 20, 12)):
        ws.column_dimensions[col].width = w
    xlsx = os.path.join(a.outdir, "02_Uebersicht_Fahrten_BahnCard.xlsx")
    wb.save(xlsx)

    # --- table as PDF (landscape), so it can be merged without Excel
    doc = fitz.open()
    W, H = fitz.paper_rect("a4-l").width, fitz.paper_rect("a4-l").height
    p = doc.new_page(width=W, height=H)
    y = 50
    p.insert_text((40, y), "Übersicht dienstlicher Fahrten mit BahnCard", fontsize=14, fontname="hebo")
    y += 28
    for label, val in (("BahnCard Nummer", number), ("BahnCard Typ", f"BahnCard {bc_type}"),
                       ("Kosten BahnCard", money(price) + " EUR"), ("Ersparnis BahnCard", money(saving) + " EUR")):
        p.insert_text((40, y), label, fontsize=10, fontname="hebo")
        p.insert_text((200, y), val, fontsize=10, fontname="helv")
        y += 16
    y += 14
    cols = [(40, "Fahrt"), (340, "Datum"), (410, "Kosten (inkl. BC)"), (510, "Tarif"), (650, "Rabatt"),
            (690, "Preis ohne BC"), (765, "Ersparnis")]
    for x, h in cols:
        p.insert_text((x, y), h, fontsize=9, fontname="hebo")
    p.draw_line((40, y + 4), (W - 30, y + 4))
    y += 18
    for r in rows:
        vals = [r[0], r[1], money(r[2]), r[3], f"{int(r[4] * 100)} %", money(r[5]), money(r[6])]
        for (x, _), v in zip(cols, vals):
            p.insert_text((x, y), str(v), fontsize=9, fontname="helv")
        y += 15
    p.draw_line((40, y - 10), (W - 30, y - 10))
    p.insert_text((690, y + 4), "Summe", fontsize=9, fontname="hebo")
    p.insert_text((765, y + 4), money(saving), fontsize=9, fontname="hebo")
    table_pdf = os.path.join(a.outdir, "02_Uebersicht_Fahrten_BahnCard.pdf")
    doc.save(table_pdf)

    # --- letter
    cc = d.get("cost_center") or prof.get("default_cost_center", "")
    name = prof.get("display_name") or prof["name"]
    body = [
        (f"Erstattung meiner MyBahnCard {bc_type} (Nummer: {number})", "hebo"),
        ("", None),
        (date, "right"),
        ("", None),
        ("Sehr geehrte Damen und Herren,", "helv"),
        ("", None),
        (f"für meine MyBahnCard {bc_type} (Nummer: {number}) habe ich am {d['purchase_date']} {money(price)} EUR "
         f"bezahlt und die im Anhang aufgelisteten dienstlich bedingten Fahrten durchgeführt.", "helv"),
        ("", None),
        (f"Der eingesparte Betrag ({money(saving)} EUR) ist höher als der Kaufpreis der BahnCard. Ich möchte Sie "
         f"deshalb bitten, mir den Betrag in Höhe von {money(price)} EUR auf mein Konto bei {prof.get('bank', '')} "
         f"(IBAN {prof.get('iban', '')}) zu überweisen.", "helv"),
        ("", None),
        (f"Die Kostenstelle zur Belastung lautet {cc}.", "helv"),
        ("", None),
        ("Freundliche Grüße", "helv"),
        ("", None), ("", None), ("", None),
        (name, "helv"),
    ]
    doc = fitz.open()
    p = doc.new_page(width=fitz.paper_rect("a4").width, height=fitz.paper_rect("a4").height)
    y = 90
    for txt, style in body:
        if not txt:
            y += 8
            continue
        if style == "right":
            p.insert_text((p.rect.width - 70 - fitz.get_text_length(txt, "helv", 11), y), txt, fontsize=11)
            y += 16
            continue
        r = fitz.Rect(70, y - 11, p.rect.width - 70, y + 200)
        rest = p.insert_textbox(r, txt, fontsize=11, fontname="hebo" if style == "hebo" else "helv")
        used = 200 - rest if rest >= 0 else 200
        y += max(16, used + 6)
    letter = os.path.join(a.outdir, "01_Erklaerung_Erstattung_BahnCard.pdf")
    doc.save(letter)

    print(f"Saving {money(saving)} EUR vs. price {money(price)} EUR -> "
          f"{'reimbursable' if saving > price else 'NOT yet reimbursable'}")
    for f in (letter, table_pdf, xlsx):
        print("Wrote", f)


if __name__ == "__main__":
    main()
